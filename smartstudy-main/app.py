import os
import time

import streamlit as st
from dotenv import load_dotenv
from pypdf import PdfReader

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import (
    GoogleGenerativeAIEmbeddings,
    ChatGoogleGenerativeAI,
)
from langchain_chroma import Chroma
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.messages import HumanMessage
from docx import Document
from pptx import Presentation
import base64


load_dotenv()

st.set_page_config(
    page_title="SmartStudy AI",
    page_icon="📚",
    layout="wide",
    initial_sidebar_state="expanded",
)




@st.cache_resource
def get_embeddings():
    return GoogleGenerativeAIEmbeddings(model="models/gemini-embedding-001")

@st.cache_resource
def get_llm():
    return ChatGoogleGenerativeAI(model="gemini-2.5-flash", temperature=0)

def get_document_text(docs):
    text = ""
    for doc in docs:
        ext = doc.name.split(".")[-1].lower()
        if ext == "pdf":
            reader = PdfReader(doc)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
        elif ext == "txt":
            text += doc.read().decode("utf-8", errors="ignore") + "\n"
        elif ext in ["docx", "doc"]:
            d = Document(doc)
            text += "\n".join([para.text for para in d.paragraphs]) + "\n"
        elif ext in ["pptx", "ppt"]:
            prs = Presentation(doc)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif ext in ["png", "jpg", "jpeg"]:
            try:
                image_bytes = doc.read()
                image_b64 = base64.b64encode(image_bytes).decode("utf-8")
                mime_type = "jpeg" if ext == "jpg" else ext
                llm = get_llm()
                msg = HumanMessage(
                    content=[
                        {"type": "text", "text": "Extract and transcribe all the text from this image. If there are tables or diagrams, describe them in detail. Return only the extracted text."},
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:image/{mime_type};base64,{image_b64}"},
                        },
                    ]
                )
                response = llm.invoke([msg])
                text += str(response.content) + "\n"
            except Exception as e:
                st.warning(f"Failed to process image {doc.name}: {e}")
    return text

def get_text_chunks(text):
    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200, length_function=len)
    return splitter.split_text(text)

def get_vector_store(chunks):
    embeddings = get_embeddings()
    return Chroma.from_texts(texts=chunks, embedding=embeddings, collection_name="smartstudy_documents")

def get_answer_with_retry(question, vector_store, retries=3, delay=3):
    for attempt in range(1, retries + 1):
        try:
            return get_answer(question, vector_store)
        except Exception as e:
            err = str(e)
            if ("503" in err or "UNAVAILABLE" in err or "wsarecv" in err or "ConnectionAborted" in err) and attempt < retries:
                time.sleep(delay * attempt)
                continue
            raise

def get_answer(question, vector_store):
    docs = vector_store.similarity_search(question, k=3)
    if not docs:
        return "I couldn't find relevant information in your uploaded documents."
    context = "\n\n".join(d.page_content for d in docs)
    prompt = ChatPromptTemplate.from_template("""
You are SmartStudy AI, a helpful and precise study assistant.
Use ONLY the document context below to answer the student's question.
Rules:
- Do not use outside knowledge.
- Do not invent facts.
- If the answer is not in the context, say: "I couldn't find this information in the uploaded documents."
- Explain clearly and directly.
- Use bullet points or numbered steps when helpful.

DOCUMENT CONTEXT:
{context}

QUESTION:
{question}

ANSWER:
""")
    llm = get_llm()
    response = (prompt | llm).invoke({"context": context, "question": question})
    if isinstance(response.content, str):
        return response.content
    if isinstance(response.content, list):
        return "\n".join(item.get("text","") for item in response.content if isinstance(item,dict) and item.get("type")=="text")
    return str(response.content)


def main():
    if "vector_store" not in st.session_state:
        st.session_state.vector_store = None
    if "messages" not in st.session_state:
        st.session_state.messages = []
    if "doc_stats" not in st.session_state:
        st.session_state.doc_stats = {"chars": 0, "chunks": 0, "docs": 0}

    # ── SIDEBAR ──
    with st.sidebar:
        st.title("📚 SmartStudy AI")
        st.caption("PDF Intelligence")

        st.subheader("📂 Upload Documents")
        uploaded_docs = st.file_uploader("Drop your documents here", accept_multiple_files=True, type=["pdf", "txt", "docx", "doc", "pptx", "ppt", "png", "jpg", "jpeg"], label_visibility="collapsed")

        if uploaded_docs:
            st.write(f"📎 {len(uploaded_docs)} file{'s' if len(uploaded_docs)>1 else ''} selected")


        process_button = st.button("⚡ Process Documents", use_container_width=True)

        if process_button:
            if not uploaded_docs:
                st.warning("Please upload at least one document first.")
            else:
                with st.spinner("🔍 Extracting and embedding your documents..."):
                    try:
                        raw_text = get_document_text(uploaded_docs)
                        if not raw_text.strip():
                            st.error("No readable text found in the uploaded documents.")
                            return
                        chunks = get_text_chunks(raw_text)
                        vs = get_vector_store(chunks)
                        st.session_state.vector_store = vs
                        st.session_state.messages = []
                        st.session_state.doc_stats = {"chars": len(raw_text), "chunks": len(chunks), "docs": len(uploaded_docs)}
                        st.success("✅ Documents processed successfully!")
                    except Exception as e:
                        st.error(f"Processing failed:\n\n{e}")



    st.title("📚 SmartStudy AI")
    st.write("Upload your study materials and have an intelligent conversation with them. Instant answers, grounded in your documents — no hallucinations.")

    # ── CHAT ──
    if not st.session_state.messages:
        st.info("Upload your documents from the sidebar, process them, then start asking questions below.")
    else:
        for message in st.session_state.messages:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

    question = st.chat_input("Ask anything about your uploaded documents...")
    if question:
        st.session_state.messages.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)

        if not st.session_state.vector_store:
            answer = "📄 Please upload and process your PDFs before asking questions."
            with st.chat_message("assistant"):
                st.markdown(answer)
        else:
            with st.chat_message("assistant"):
                with st.spinner("🔎 Searching your documents..."):
                    try:
                        answer = get_answer_with_retry(question, st.session_state.vector_store)
                    except Exception as e:
                        answer = f"⚠️ I encountered an error:\n\n{e}"
                st.markdown(answer)

        st.session_state.messages.append({"role": "assistant", "content": answer})


if __name__ == "__main__":
    main()
